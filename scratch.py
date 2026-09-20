import docx
import sys
import glob

files = glob.glob('Proyek/*.docx')
for f in files:
    print('---', f, '---')
    doc = docx.Document(f)
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f"[{i}] {p.text}")
    print("\nTables:")
    for t in doc.tables:
        for r in t.rows:
            row_data = [c.text.replace('\n', ' ') for c in r.cells]
            print(row_data)
